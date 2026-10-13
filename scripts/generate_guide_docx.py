import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page Margins (0.75 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Colors
    NAVY = RGBColor(0x0F, 0x25, 0x37)
    BLUE = RGBColor(0x1A, 0x73, 0xE8)
    MUTED = RGBColor(0x55, 0x55, 0x55)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(26)
        run.bold = True
        run.font.color.rgb = NAVY

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.italic = True
        run.font.color.rgb = MUTED

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.bold = True
        run.font.color.rgb = NAVY

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = BLUE

    def add_p(text, bold_prefix=None, italic=False, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = NAVY
        run = p.add_run(text)
        run.italic = italic
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = NAVY
        p.add_run(text)
        return p

    def add_callout(text, title="KEY TAKEAWAY", fill_hex="E8F0FE"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(7.0)
        set_cell_background(cell, fill_hex)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        r_title = p.add_run(f"[{title}] ")
        r_title.bold = True
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = BLUE

        r_text = p.add_run(text)
        r_text.font.size = Pt(10.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_styled_table(headers, rows):
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col_idx, h in enumerate(headers):
            cell = table.cell(0, col_idx)
            set_cell_background(cell, "0F2537")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(h)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(10)

        for row_idx, row_data in enumerate(rows, start=1):
            for col_idx, text in enumerate(row_data):
                cell = table.cell(row_idx, col_idx)
                set_cell_background(cell, "F8F9FA" if row_idx % 2 == 1 else "FFFFFF")
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.12
                r = p.add_run(str(text))
                r.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_image_box(image_path, caption):
        full_path = os.path.join(r"e:\My Projects\ai-buildercup-2026", image_path)
        if not os.path.exists(full_path):
            print(f"Warning: image not found: {full_path}")
            return
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        run_img = p_img.add_run()
        run_img.add_picture(full_path, width=Inches(6.5))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption)
        r_cap.font.size = Pt(9.5)
        r_cap.italic = True
        r_cap.font.color.rgb = MUTED

    # ------------------ DOCUMENT CONTENT ------------------

    add_title("CatSight — Business & Domain Guide")
    add_subtitle(
        "A Plain-English Explanation of Reinsurance Math, Catastrophe Claims, and AI Contract Intelligence\n"
        "Team Techno Crackers: Venket Raj S · Malik Chandra Pandian · Arun Veerappan · Nagaraj\n"
        "Live URL: https://catsight-api-342162443546.us-central1.run.app | Code: https://github.com/venketraj/catsight-ai-buildercup-2026\n"
        "Demo Video (< 3 min): [VIDEO_LINK_PLACEHOLDER — Insert Public YouTube / Demo Video URL (< 3:00)]"
    )

    add_callout(
        "This guide translates the entire CatSight project (synced to the final submission build, 874 automated tests, "
        "and 14-slide official AI Builder Cup submission deck) into everyday business language. It demystifies every reinsurance formula, "
        "explains why primary insurers and reinsurers argue over millions of dollars after natural disasters, and shows "
        "how CatSight delivers a verified portfolio view in 0.7–1.0 seconds (math) and ~10–12 seconds p50 (full AI wording check).",
        title="DOCUMENT PURPOSE"
    )

    # ------------------ SECTION 1 ------------------
    add_h1("1. The Real-World Problem: The Morning After a Disaster")
    add_p(
        "Imagine a Category 5 typhoon—like Typhoon Jebi in September 2018—slams into Osaka, Japan. In reality, Jebi produced "
        "857,284 insurance claims and ¥1,067.8bn paid out (the largest ever for a Japanese typhoon, per GIAJ). Early modelled "
        "estimates 10 days after landfall were $3–5.5bn (Moody's RMS), but the final insured loss grew to ~$13bn (Swiss Re)."
    )
    add_p(
        "Who pays for all of this? Local primary insurance companies (in our app, fictional insurers like \"Sakura General Insurance\" "
        "and \"Yamabiko Casualty\") sell property insurance to households and factories. When a typhoon strikes, thousands of customers "
        "file claims simultaneously. If a local insurer had to pay hundreds of millions of dollars from its own bank account, it would go bankrupt."
    )
    add_p(
        "To protect themselves, primary insurers buy \"Insurance for Insurance Companies\"—called REINSURANCE—from global reinsurers "
        "(like Swiss Re, Munich Re, or in our app, our fictional company \"Lion Re\")."
    )

    add_image_box("scratch/slides/scene1_hook_slide.png", "Figure 1: The morning after a typhoon — why reinsurers need a fast, defensible view of their own book.")

    add_p(
        "The Bottleneck Today: The morning after a typhoon, the CFO of Lion Re asks: \"Which of our treaties are hit, how much do we owe, "
        "and which contract clauses could change that number?\""
    )
    add_styled_table(
        ["Source / Workflow", "When It Arrives", "Why It Falls Short for One Reinsurer"],
        [
            ["Modeller event response (Moody's RMS)", "4 to 20 days", "Describes the overall market or pre-coded terms, not your PDF wording"],
            ["PERILS (Japan typhoon & flood)", "6 weeks", "Market-wide industry loss report, not treaty-specific"],
            ["Verisk PCS Japan (events >= $2bn)", "Up to 90 days", "Market-wide catastrophe index"],
            ["In-house manual analyst workflow", "2 to 4 business days (30–60 analyst-hours)", "Spreadsheets + re-reading 25+ exposed PDF contracts by hand"],
            ["CatSight (on the reinsurer's own book)", "0.7–1.0s (math) · ~10–12s p50 (verified AI flags)", "Joins event × exposure × layers × verified PDF wording"]
        ]
    )
    add_p(
        "Industry context (sourced): 48% of insurers do not license catastrophe models, and nearly 60% run cat teams of five or fewer people (Aon 2025). "
        "Moody's notes that outward reinsurance treaty structures often \"only exist in their capital model, a homegrown tool, or even in a spreadsheet.\" "
        "Meanwhile, Singapore writes S$27.6bn of reinsurance premium (hosting 16 of the top 25 global reinsurers' regional hubs), yet 91% of Asia-Pacific's "
        "2023 catastrophe losses ($65bn) were uninsured (MAS)."
    )

    add_callout(
        "Where CatSight Fits (Complementary, Not a Replacement): CatSight is NOT a catastrophe model and does not compete with Moody's RMS or Verisk. "
        "Cat models estimate HOW BIG the loss is. CatSight is the layer after that: it joins an event footprint (public NOAA/USGS data today, or a "
        "licensed vendor footprint via the same typed interface) to the reinsurer's own treaties and reads their PDF wording.",
        title="POSITIONING"
    )

    # ------------------ SECTION 2 ------------------
    add_h1("2. Demystifying the Reinsurance Jargon (In Plain English)")
    add_p("You do not need an actuarial background to understand CatSight. Here are the six core terms:")

    add_h2("A. Cedent vs. Reinsurer")
    add_bullet(
        "The primary insurance company that sells policies to homeowners and businesses (e.g., Sakura General Insurance, Yamabiko Casualty, "
        "Kogane Fire & Marine). They are called the 'Cedent' because they 'cede' (pass on) part of their disaster risk.",
        bold_prefix="The Cedent (Primary Insurer): "
    )
    add_bullet(
        "The company that accepts that ceded risk in exchange for an annual premium. In CatSight, we are Lion Re, holding a book of 24 treaties "
        "across 8 cedents in Japan, the Philippines, and Taiwan.",
        bold_prefix="The Reinsurer (Lion Re — \"Us\"): "
    )

    add_h2("B. What is a \"Treaty\"? (Cat XL vs. Quota Share)")
    add_p("A Treaty is simply a formal 1-year contract PDF between a Cedent and its Reinsurers. CatSight handles two types of treaties:")
    add_bullet(
        "Works like stacked deductibles (layers). Reinsurers only pay when losses exceed a set threshold (Retention), up to a cap (Limit).",
        bold_prefix="1. Catastrophe Excess of Loss (CAT_XL): "
    )
    add_bullet(
        "A simpler proportional split: the reinsurer takes a fixed percentage (e.g., 40% cession_pct) of every dollar of loss from the very first dollar, "
        "sometimes up to an event_limit (e.g., Treaty T-006, T-021, T-022, T-023).",
        bold_prefix="2. Quota Share (QS): "
    )

    add_h2("C. Layers, Retention, and Limit (\"$20M xs $10M\")")
    add_p(
        "In a Cat XL treaty, coverage is sliced into horizontal tiers called Layers. When you see \"$20M xs $10M\", read it as "
        "\"$20 Million limit in excess of a $10 Million retention (deductible)\":"
    )
    add_styled_table(
        ["Tier / Layer", "Contract Shorthand", "Plain-English Meaning", "Who Pays?"],
        [
            ["Cedent Retention", "First $10M ($0 to $10M)", "The insurer's deductible. The primary insurer pays the first $10M.", "Cedent (100%)"],
            ["Layer 1 (L1)", "$20M xs $10M", "Covers losses above $10M, up to a max payout of $20M (i.e., $10M to $30M).", "Reinsurers on L1"],
            ["Layer 2 (L2)", "$30M xs $30M", "Covers losses above $30M, up to a max payout of $30M (i.e., $30M to $60M).", "Reinsurers on L2"],
            ["Layer 3 (L3)", "$40M xs $60M", "Covers losses above $60M, up to a max payout of $40M (i.e., $60M to $100M).", "Reinsurers on L3"]
        ]
    )

    add_h2("D. \"Our Share\" vs. \"All Reinsurers\"")
    add_p(
        "A single reinsurer rarely takes 100% of a layer. Instead, a syndicate of reinsurers each subscribes to a percentage share. "
        "For example, on Sakura's Layer 1 ($20M limit), Lion Re has a 25% share ($5M max payout), and on Layer 2 ($30M limit), Lion Re has a 10% share. "
        "That is why CatSight always shows three party views: What the Insurer Keeps, What All Reinsurers Pay, and Lion Re's Net (\"Our Net\")."
    )

    add_h2("E. Burn (Layer Erosion)")
    add_p(
        "\"Burn\" is simply the percentage of a layer's limit used up by the disaster. "
        "If Layer 1 has a $20M limit and pays $20M, Burn = 100% (EXHAUSTED). If Layer 2 has a $30M limit and pays $24.3M, Burn = 81.0% ($24.3M / $30M)."
    )

    add_h2("F. Reinstatements & Reinstatement Premium (RIP) — Reloading the Shield")
    add_p(
        "Think of a reinsurance layer as a protective shield. When Typhoon Jebi burns 100% of Layer 1 in September, that shield is used up. "
        "If another typhoon (like Typhoon Trami) hits three weeks later, the primary insurer would have zero Layer 1 protection left!"
    )
    add_p(
        "To prevent this, treaties include a REINSTATEMENT clause: the reinsurers automatically reload the burned portion of the layer for the rest "
        "of the contract year. However, reloading is not free—the primary insurer must pay a reload fee called a REINSTATEMENT PREMIUM (RIP):"
    )
    add_callout(
        "Exact Reinstatement Premium (RIP) Rule (Per Occurrence):\n"
        "RIP = (Limit Eroded in Occurrence / Layer Limit) × Layer Premium × Reinstatement Rate\n\n"
        "Key Financial Effect:\n"
        "• Insurer's Net Cost = Retention + RIP Payable (+ any Excluded Loss)\n"
        "• Reinsurers' Net Cost = Claims Paid (Ceded) - RIP Received\n"
        "Because the cedent pays RIP to the reinsurer, RIP increases the insurer's cost and reduces the reinsurer's net loss!",
        title="HOW RIP AFFECTS THE MONEY"
    )

    # ------------------ SECTION 3 ------------------
    add_h1("3. Step-by-Step Numerical Examples (How the Engine Calculates)")
    add_p(
        "Every money figure in CatSight comes from deterministic Python using exact Decimal math (874 automated tests: 861 unit + 13 live). "
        "Here are the two core worked examples from our documentation (`docs/DOMAIN_PRIMER.md` and `docs/MONEY_MOMENT.md`):"
    )

    add_h2("Example 1: One Event vs. Two Events Under an Hours Clause ($54.4M Gross Loss)")
    add_p(
        "Suppose Sakura General Insurance suffers $54.4M in gross losses. Its treaty has Layer 1 ($20M xs $10M, premium $2.0M, Lion Re share 25%, "
        "1 reinstatement @ 100%) and Layer 2 ($30M xs $30M, premium $1.5M, Lion Re share 10%)."
    )
    add_bullet(
        "First $10.0M is Cedent Retention. Layer 1 pays its full $20.0M limit (100% burn). Layer 2 pays $24.4M ($54.4M - $30M, 81.3% burn). "
        "Total Ceded = $44.4M. Lion Re gross loss = (25% × $20M) + (10% × $24.4M) = $5.00M + $2.44M = $7.44M. "
        "Lion Re RIP received = (20/20 × $2.0M × 25%) + (24.4/30 × $1.5M × 10%) = $0.500M + $0.122M = $0.622M. "
        "Lion Re Net = $7.44M - $0.622M = $6.82M.",
        bold_prefix="Reading A — Treated as ONE Event ($54.4M): "
    )
    add_bullet(
        "If the hours clause splits the storm into two separate occurrences, the insurer must pay its $10M deductible TWICE ($20.0M total retention)! "
        "Occurrence 1 ($35.0M) burns $20.0M of L1 and $5.0M of L2. Occurrence 2 ($19.4M) burns $9.4M of the reinstated L1 and $0 of L2. "
        "Total Ceded drops from $44.4M to $34.4M (-$10.0M / -22.5%). Layer 2 drops by -80% ($24.4M -> $5.0M). "
        "Wait—what happens to Lion Re? Because Lion Re holds 25% of L1 (which now pays $29.4M instead of $20.0M) and only 10% of L2, "
        "Lion Re's net loss actually INCREASES from $6.82M to $7.33M (+$0.51M / +7.4%)!",
        bold_prefix="Reading B — Split into TWO Events ($35.0M + $19.4M): "
    )
    add_callout(
        "Why This Matters: In Example 1, splitting the storm into two events saves the overall reinsurance market $9.03M net, "
        "but COSTS Lion Re $0.51M more because Lion Re's share is concentrated in Layer 1! A single headline number would hide that.",
        title="THE ASYMMETRY INSIGHT"
    )

    add_h2("Example 2: The Live Jebi Clause Courtroom — Storm-Surge Exclusion (Treaty T-004)")
    add_p(
        "On Jebi's real 9-hour timeline, all wind damage occurs within 9 hours (so no 72-hour clause is breached in reality; hours-clause splits "
        "are explored in the What-if panel, clearly labelled HYPOTHETICAL). For Jebi, the real live dispute is on Treaty T-004 (Yamabiko Casualty): "
        "JMA recorded a 329 cm storm surge in Osaka Bay that flooded Kansai Airport, and T-004 (Page 2, Clause 8) excludes Storm Surge!"
    )
    add_styled_table(
        ["Party / Metric", "Base (0% Excluded)", "20% Surge Share (Recorded Live Run)", "30% Surge Share (Reference)", "Direction"],
        [
            ["Gross Loss to Treaty T-004", "$58.18M", "$46.54M (-$11.64M)", "$40.72M (-$17.46M)", "Excluded loss leaves the treaty"],
            ["Total Ceded (Gross Recoveries)", "$39.58M", "$27.94M (-$11.64M)", "$22.12M (-$17.46M)", "Less loss reaches Layer 1 & Layer 2"],
            ["Insurer's Net Cost (Retention + RIP + Excluded)", "$22.79M", "$33.10M (+$10.31M)", "$38.30M (+$15.51M)", "Insurer carries the excluded surge loss"],
            ["All Reinsurers Pay (Recoveries - RIP)", "$35.39M", "$25.08M (-$10.31M)", "$19.88M (-$15.51M)", "Reinsurers save exactly what insurer loses"],
            ["Lion Re Net (\"Our Net\")", "$7.61M", "$6.27M (-$1.34M)", "$5.17M (-$2.44M)", "Lion Re saves money under exclusion"]
        ]
    )

    # ------------------ SECTION 4 ------------------
    add_h1("4. The CatSight 5-Step Analyst Walkthrough (With Screenshots)")
    add_p(
        "The live application (https://catsight-api-342162443546.us-central1.run.app) is laid out for a judge or analyst in five clear steps "
        "under a sticky summary bar. Visitors can click 'Instant replay' (~6s recorded real run) or 'Run live' (~15s live Gemini agents):"
    )

    add_h2("Step 1: Event — The Disaster Facts")
    add_p(
        "Loads the historical catastrophe track (NOAA IBTrACS for typhoons or USGS ShakeMap for earthquakes), event window (start, end, duration_h), "
        "and verified JMA facts (e.g., Kansai Airport 58.1 m/s gust; Osaka 329 cm record storm-surge tide)."
    )
    add_image_box("scratch/step1_event.png", "Figure 2: Step 1 (Event) — Selecting Typhoon Jebi (2018) and viewing its storm track and JMA facts.")

    add_h2("Step 2: Who's Hit — Regional Exposure Overlay")
    add_p(
        "Colours all 34 affected Japanese prefectures by peak wind or shaking intensity. Clicking any region reveals the insured property value "
        "by insurer, alongside a full accessible table alternative."
    )
    add_image_box("scratch/step2_hit.png", "Figure 3: Step 2 (Who's Hit) — Choropleth map of peak winds and regional insured values.")

    add_h2("Step 3: What It Costs — Deterministic Layer Waterfall")
    add_p(
        "Within 0.7–1.0 seconds, the deterministic loss engine matches all 24 treaties (13 matched, 11 excluded with reasons—e.g., PHL or TWN territory) "
        "and computes gross loss ($733.04M), insurer retention ($242.12M), all reinsurers ($490.92M), and Lion Re's net ($70.36M). "
        "Clicking 'How was this calculated?' shows every treaty's region-by-region damage and layer formulas."
    )
    add_image_box("scratch/step3_cost.png", "Figure 4: Step 3 (What It Costs) — Three-party financial bar and treaty-by-treaty layer breakdown.")

    add_h2("Step 4: Clauses That Matter — Verified Citations & The Endorsement Detective")
    add_p(
        "While the analyst views the numbers, Gemini 3.5 Flash-Lite searches the 252 clause vectors in Firestore and raises wording flags. "
        "Every candidate flag passes through our Verifier Loop: (1) deterministic normalized & word-aligned '...' fragment match against the stored clause text, and "
        "(2) a semantic support check on Gemini 3.8 Flash. Failed flags get one revision retry; if they still fail, they are WITHHELD and counted "
        "in the header ('n of n displayed citations verified · k withheld')."
    )
    add_image_box("scratch/step4_clauses.png", "Figure 5: Step 4 (Clauses That Matter) — Machine-verified wording flags and the bounded What-if tool.")

    add_p(
        "The Endorsement Detective: Treaty T-004 has a 72-hour windstorm clause on Page 2, but Page 7 is a scanned, stamped image endorsement "
        "with no text layer that amends Windstorm to 168 consecutive hours. Gemini transcribes the image once at ingest, flags the override, "
        "and badges citations to it distinctly as 'transcribed' rather than 'verified'."
    )
    add_image_box("scratch/step4_pdf_viewer.png", "Figure 6: The Treaty Clause Viewer showing T-004's transcribed scanned endorsement overriding the base contract.")

    add_h2("Step 4b: The Clause Courtroom (Adversarial Multi-Agent Fan-Out)")
    add_p(
        "When a verified HIGH or MEDIUM exclusion flag is present (like T-004's storm-surge exclusion), the analyst can open the Clause Courtroom. "
        "Two ADK agents (cedent_counsel and reinsurer_counsel) run in parallel, blind to each other, arguing opposing readings from verified quotes "
        "(with 1 revision retry and '...' fragment matching). Each states an excluded surge share assumption (bounded 0–50%), the Python engine prices both sides, "
        "and an Arbiter (on Gemini 3.8 Flash) summarises the dispute without deciding it."
    )
    add_image_box("scratch/step4_courtroom_resolved.png", "Figure 7: The Clause Courtroom — Parallel AI counsel, engine-priced outcomes, and analyst position recording.")

    add_h2("Step 5: Decide & Export")
    add_p(
        "The analyst records their position on the dispute (Insurer's reading, Reinsurers' reading, or Unresolved) with an audit timestamp, "
        "reviews the AI-generated report (where every single number is checked in Python against the engine's output), and downloads a "
        "formatted Word (.docx) or Markdown underwriting memo."
    )
    add_image_box("scratch/step5_decide.png", "Figure 8: Step 5 (Decide) — Traceable executive report and one-click Word (.docx) / Markdown export.")

    # ------------------ SECTION 5 ------------------
    add_h1("5. Google Cloud Architecture & Measured Proof Pack")
    add_p(
        "CatSight runs as a single Cloud Run service serving both the FastAPI SSE backend and the React UI from the same origin, "
        "backed by two Google ADK 2.11 Workflow graphs, Gemini Enterprise Agent Platform (formerly Vertex AI) global endpoints, Firestore Native vector search, and Cloud Trace/Monitoring."
    )
    add_image_box("scratch/slides/scene6_arch_slide.png", "Figure 9: Google Cloud Architecture — ADK 2.11 Workflows, Gemini Enterprise Agent Platform (Vertex AI), Firestore Vector Search, and Cloud Run.")

    add_styled_table(
        ["Benchmark / Metric", "Measured Result", "Notes & Audit Trail (`docs/RESULTS.md`)"],
        [
            ["Loss engine vs. hand-checked worked example", "Exact (100%)", "874 automated tests (`Decimal` half-up; 861 unit + 13 live)"],
            ["Time to money on screen (two 5-run live benchmarks)", "0.7–1.0 s p50 (0.6–1.1 s)", "Engine runs before LLM agents"],
            ["Complete live analysis with verified flags", "9.8–12.1 s p50 (8.0–31.2 s)", "Warm Cloud Run service, SSE with 15s heartbeat"],
            ["Model cost per live analysis", "$0.0094–$0.0102 p50 ($0.0081–$0.0126)", "Measured from Gemini token usage"],
            ["Clause Courtroom latency & cost", "20.3 s · $0.0040", "Parallel counsel on 3.5 Flash-Lite, Arbiter on 3.8 Flash"],
            ["Treaty-term extraction (RAG-first, 8 wordings)", "99.5% (185/186 fields) in 57 s", "Upper bound on tuned set (vs ~25 min whole-doc)"],
            ["Governing clause retrieved", "186/186 (100%)", "Firestore 768-dim vector search with treaty_id pre-filter"],
            ["Per-field citation audit", "186/188 cited · 185/186 verified", "184 verified + 1 transcribed; 1 unverified withheld"],
            ["Untraceable figures in AI reports", "0 across all live runs", "Every number checked against engine output in code"]
        ]
    )

    output_path = r"e:\My Projects\ai-buildercup-2026\docs\CatSight_Business_and_Domain_Guide.docx"
    try:
        doc.save(output_path)
        print(f"Updated guide successfully saved to: {output_path}")
    except PermissionError:
        alt_path = r"e:\My Projects\ai-buildercup-2026\docs\CatSight_Business_and_Domain_Guide_v2.docx"
        doc.save(alt_path)
        print(f"Original file is open in Word; saved updated guide to: {alt_path}")

if __name__ == "__main__":
    create_document()

