"""FR-EXPORT - the analyst report as a Word document.

Built from two sources, each authoritative for its part:
  * the engine's figures, recomputed at export time (deterministic, so they
    are the figures the run showed and cannot drift from the computation);
  * the stored run - verified flags with citations, what was withheld, the
    report text, the cost, and whether it was a replay.
Nothing here computes money; figures are formatted as served.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any, Mapping

__all__ = ["usd", "build_report_docx"]

DISCLAIMER = (
    "Indicative first view on synthetic data - not a catastrophe-model output and not "
    "advice. Figures are deterministic loss-engine estimates; wording flags are "
    "machine-verified citations for an analyst to review. Synthetic portfolio and "
    "wordings; cedent names are fictional. Hazard: NOAA IBTrACS / USGS ShakeMap."
)


def usd(value_in_millions: Any) -> str:
    """'54.30' (USD millions, as served) -> '$54.30M'; '41888.0' -> '$41.89B'."""
    if value_in_millions is None:
        return "-"
    amount = Decimal(str(value_in_millions))
    if abs(amount) >= 1000:
        return f"${(amount / 1000).quantize(Decimal('0.01'))}B"
    return f"${amount.quantize(Decimal('0.01'))}M"


def _table(doc: Any, header: list[str], rows: list[list[str]], numeric: set[int] = frozenset()) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, title in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = title
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            if i in numeric:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


_BOLD = re.compile(r"(\*\*[^*]+\*\*)")


def _rich(paragraph: Any, text: str) -> None:
    for part in _BOLD.split(text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def _markdown(doc: Any, text: str) -> None:
    """The reporter's small Markdown: ## headings, bullets, bold."""
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("#"):
            doc.add_heading(line.lstrip("#").strip(), level=3)
            continue
        bullet = re.match(r"^\s*[*-]\s+(.*)$", line)
        if bullet:
            _rich(doc.add_paragraph(style="List Bullet"), bullet.group(1))
        else:
            _rich(doc.add_paragraph(), line)


def build_report_docx(
    run: Mapping[str, Any], impact: Mapping[str, Any], *, generated_at: datetime | None = None
) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    generated_at = generated_at or datetime.now(timezone.utc)
    t = impact["totals"]
    doc = Document()
    doc.core_properties.title = f"CatSight analysis - {impact['event']['name']}"
    doc.core_properties.author = "CatSight"

    doc.add_heading(f"CatSight analysis - {impact['event']['name']}", level=0)
    meta = doc.add_paragraph()
    meta.add_run(
        f"Generated {generated_at.strftime('%Y-%m-%d %H:%M')} UTC  ·  perspective: Lion Re (reinsurer)  ·  "
        f"basis: {'as-if on the current book' if impact.get('as_if') else 'actual'}  ·  "
        f"method {impact.get('methodology', {}).get('version') or 'n/a'}  ·  run {run['run_id']}"
    ).italic = True
    if run.get("replay"):
        note = doc.add_paragraph()
        note.add_run(f"Replay: {run['replay'].get('message')}").bold = True

    doc.add_heading("Headline figures", level=1)
    _table(doc, ["", "Amount"], [
        ["Insurers' claims (gross)", usd(t["gross"])],
        ["Insurers keep", usd(t.get("retained"))],
        ["Reinsurers pay (ceded)", usd(t["ceded"])],
        ["Our net (Lion Re)", usd(t["our_net"])],
        ["Treaties responding / not", f"{t['treaties_matched']} / {t['treaties_excluded']}"],
    ], numeric={1})

    doc.add_heading("Treaties that responded", level=1)
    _table(doc, ["Treaty", "Insurer", "Type", "Gross", "Ceded", "Our net"], [
        [r["treaty_id"], r["cedent"], r["type"], usd(r["gross"]), usd(r["ceded"]), usd(r["our_net"])]
        for r in impact["treaties"]
    ], numeric={3, 4, 5})

    doc.add_heading("Wording to review", level=1)
    doc.add_paragraph(run.get("header") or "The wording check did not run.").italic = True
    if not run.get("verified") and run.get("header"):
        doc.add_paragraph("No clause in the matched wordings changes the answer for this event.")
    for item in run.get("verified", []):
        flag = item["flag"]
        heading = doc.add_paragraph()
        heading.add_run(f"{flag['severity']}  ·  {flag['treaty_id']}  ·  {flag['issue_type']}").bold = True
        doc.add_paragraph(flag["explanation"])
        quote = doc.add_paragraph(f"“{flag['citation']['quote']}”", style="Intense Quote")
        quote.paragraph_format.space_after = Pt(2)
        badge = doc.add_paragraph().add_run(item.get("badge") or "verified")
        badge.font.size = Pt(9)
        badge.font.color.rgb = RGBColor(0x1C, 0x7C, 0x4A)
    if run.get("withheld"):
        doc.add_paragraph(f"Withheld - failed verification, not shown as findings: {len(run['withheld'])}")
        for w in run["withheld"]:
            doc.add_paragraph(f"{w.get('treaty_id')} {w.get('issue_type')}: {w.get('reason') or 'no reason recorded'}",
                              style="List Bullet")

    doc.add_heading("Analyst report", level=1)
    _markdown(doc, run.get("report") or "")

    if impact.get("excluded"):
        doc.add_heading("Did not respond", level=1)
        for e in impact["excluded"]:
            doc.add_paragraph(f"{e['treaty_id']}: {e['reason']}", style="List Bullet")

    doc.add_heading("Notes", level=1)
    doc.add_paragraph(DISCLAIMER)
    if run.get("cost_usd"):
        doc.add_paragraph(f"Model cost of this run: ${run['cost_usd']}.")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "CatSight · indicative first view on synthetic data · not advice"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
