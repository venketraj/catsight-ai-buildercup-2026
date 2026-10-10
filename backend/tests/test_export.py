"""Stored runs (app.runs) and the Word export (app.report_docx, FR-EXPORT)."""

from __future__ import annotations

import asyncio
import io
import json
from types import SimpleNamespace

from fastapi.testclient import TestClient

from app.analysis import analyse
from app.config import load_settings
from app.main import create_app
from app.report_docx import build_report_docx, usd
from app.runs import FirestoreRunStore, MemoryRunStore, capture, new_run_id
from catsight_agent.store import LocalStore

STORE = LocalStore("data", today="2026-10-10")
FLAG = {"treaty_id": "T-004", "issue_type": "EXCLUSION", "severity": "HIGH",
        "explanation": "Storm surge is excluded and the event caused a record tide.",
        "citation": {"treaty_id": "T-004", "page": 2, "clause_no": "8", "source": "text",
                     "quote": "storm surge, howsoever caused"}}


def _sse(frame):
    return f"data: {json.dumps(frame)}\n\n"


async def _gen(frames):
    for f in frames:
        yield _sse(f)


def _collect(agen):
    async def go():
        return [c async for c in agen]

    return asyncio.run(go())


COMPLETE = [
    {"type": "notice", "kind": "replay", "message": "Instant replay of a real live run",
     "recorded_at": "2026-10-10T08:31:21+00:00"},
    {"type": "flag", "status": "candidate", "flag": FLAG},
    {"type": "flag", "status": "verified", "badge": "verified · T-004, p. 2, cl. 8", "flag": FLAG},
    {"type": "verification", "header": "1 of 1 displayed citations verified · 1 withheld",
     "verified": 1, "withheld": [{"treaty_id": "T-002", "issue_type": "EXCLUSION", "reason": "carve-back"}]},
    {"type": "report", "markdown": "## Summary\nOur net is **70.36**.\n- one\n\n# Top", "cost_usd": "0.0081",
     "unsupported_figures": [], "calls": {}},
    {"type": "done", "analysis_id": "replay-jebi-2018", "agents": True},
]


# --------------------------------------------------------------------------- #
# Runs
# --------------------------------------------------------------------------- #


def test_a_completed_stream_is_saved_and_the_done_frame_names_it():
    store = MemoryRunStore()
    out = _collect(capture(_gen(COMPLETE + []), store, "run-1", "jebi-2018"))
    done = json.loads(out[-1][5:])
    assert done["run_id"] == "run-1" and done["agents"] is True
    run = store.get("run-1")
    assert run["replay"]["recorded_at"].startswith("2026-10-10")
    assert [v["flag"]["treaty_id"] for v in run["verified"]] == ["T-004"]
    assert run["withheld"][0]["reason"] == "carve-back"
    assert run["cost_usd"] == "0.0081" and run["completed_at"].endswith("+00:00")


def test_a_stream_without_a_report_is_not_saved():
    store = MemoryRunStore()

    async def frames():
        yield "not a frame\n\n"
        yield "data: {broken\n\n"
        yield _sse({"type": "done", "analysis_id": "x"})

    out = _collect(capture(frames(), store, "run-2", "x"))
    assert store.get("run-2") is None
    assert "run_id" not in json.loads(out[-1][5:]) and len(out) == 3


def test_a_failed_save_does_not_break_the_stream():
    class Broken(MemoryRunStore):
        def save(self, run_id, run):
            raise RuntimeError("firestore down")

    out = _collect(capture(_gen(COMPLETE), Broken(), "run-3", "jebi-2018"))
    assert "run_id" not in json.loads(out[-1][5:])


def test_memory_store_is_bounded_and_ids_are_unique():
    store = MemoryRunStore(limit=2)
    for i in range(3):
        store.save(f"r{i}", {"i": i})
    assert store.get("r0") is None and store.get("r2") == {"i": 2}
    assert new_run_id() != new_run_id() and new_run_id().startswith("run-")


def test_firestore_run_store():
    docs = {}

    class Doc:
        def __init__(self, key):
            self.key = key

        def set(self, data):
            docs[self.key] = data

        def get(self):
            return SimpleNamespace(exists=self.key in docs, to_dict=lambda: docs.get(self.key))

    client = SimpleNamespace(collection=lambda name: SimpleNamespace(document=Doc))
    store = FirestoreRunStore(client)
    store.save("run-9", {"a": 1})
    assert store.get("run-9") == {"a": 1} and store.get("nope") is None


# --------------------------------------------------------------------------- #
# Word document
# --------------------------------------------------------------------------- #


def test_usd_formatting():
    assert usd("54.30") == "$54.30M" and usd("41888.0") == "$41.89B" and usd(None) == "-"


def _text(docx_bytes):
    from docx import Document

    doc = Document(io.BytesIO(docx_bytes))
    body = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    return doc, body, cells


def test_the_word_report_carries_figures_flags_citations_and_disclaimer():
    store = MemoryRunStore()
    _collect(capture(_gen(COMPLETE), store, "run-1", "jebi-2018"))
    impact = analyse("jebi-2018", store=STORE).to_json()
    doc, body, cells = _text(build_report_docx(store.get("run-1"), impact))

    assert doc.core_properties.title == "CatSight analysis - Typhoon Jebi"
    assert "perspective: Lion Re (reinsurer)" in body and "method 2026.1" in body
    assert "Replay: Instant replay of a real live run" in body
    assert "$70.36M" in cells and "$733.04M" in cells and "T-001" in cells
    assert "HIGH  ·  T-004  ·  EXCLUSION" in body and "storm surge, howsoever caused" in body
    assert "verified · T-004, p. 2, cl. 8" in body and "carve-back" in body
    assert "Summary" in body and "Our net is 70.36." in body
    assert "not a catastrophe-model output" in body and "$0.0081" in body
    assert "PHL" in body  # a "did not respond" reason


def test_a_run_without_flags_says_so():
    run = {"run_id": "run-x", "event_id": "noto-2024", "verified": [], "withheld": [],
           "header": "0 of 0 displayed citations verified · 0 withheld", "report": "Plain text.",
           "replay": None, "cost_usd": None}
    impact = analyse("noto-2024", store=STORE).to_json()
    _, body, _ = _text(build_report_docx(run, impact))
    assert "No clause in the matched wordings changes the answer" in body
    empty = dict(run, header=None, report=None)
    _, body, _ = _text(build_report_docx(empty, dict(impact, excluded=[])))
    assert "The wording check did not run." in body


# --------------------------------------------------------------------------- #
# API
# --------------------------------------------------------------------------- #


def test_download_endpoint(tmp_path):
    from app.replay import fixture_path

    fixture_path(tmp_path, "jebi-2018").write_text(json.dumps({
        "event_id": "jebi-2018", "recorded_at": "2026-10-10T08:31:21+00:00", "seconds": 0.1,
        "frames": [{"t": 0, "frame": f} for f in COMPLETE[1:]],
    }), encoding="utf-8")
    api = create_app(load_settings({"GOOGLE_CLOUD_PROJECT": "p", "REPLAY_FIXTURE_DIR": str(tmp_path)}),
                     store=STORE)
    client = TestClient(api)
    stream = client.post("/api/analyses", json={"event_id": "jebi-2018", "replay": True}).text
    done = [json.loads(l[5:]) for l in stream.split("\n\n") if l.startswith("data:")][-1]

    response = client.get(f"/api/runs/{done['run_id']}/report.docx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f"catsight-jebi-2018-{done['run_id']}.docx" in response.headers["content-disposition"]
    _, body, _ = _text(response.content)
    assert "Typhoon Jebi" in body
    assert client.get("/api/runs/run-nope/report.docx").status_code == 404
